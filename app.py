import os
import io
import json
import re
import base64
import difflib
import anthropic
from flask import Flask, request, jsonify, render_template
from docx import Document

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50MB

client = anthropic.Anthropic()

CHUNK_SIZE = 30  # paragraphs per Claude API call


def get_para_text(para):
    return "".join(run.text for run in para.runs)


def apply_correction(para, corrected_text):
    """Apply corrected text to paragraph while preserving run-level formatting."""
    original_text = get_para_text(para)
    if not original_text.strip() or original_text == corrected_text:
        return False
    runs = para.runs
    if not runs:
        return False

    # Map each character in original_text to its run index
    char_run_map = []
    for run_idx, run in enumerate(runs):
        char_run_map.extend([run_idx] * len(run.text))

    if not char_run_map:
        return False

    matcher = difflib.SequenceMatcher(None, original_text, corrected_text, autojunk=False)
    new_run_texts = [""] * len(runs)

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for k in range(i2 - i1):
                run_idx = char_run_map[i1 + k]
                new_run_texts[run_idx] += corrected_text[j1 + k]
        elif tag in ("replace", "insert"):
            # Assign new characters to the run at the start of the changed region
            if i1 < len(char_run_map):
                run_idx = char_run_map[i1]
            elif i1 > 0:
                run_idx = char_run_map[i1 - 1]
            else:
                run_idx = 0
            new_run_texts[run_idx] += corrected_text[j1:j2]
        # "delete": original chars are removed, nothing to add

    for run, new_text in zip(runs, new_run_texts):
        run.text = new_text
    return True


def collect_paragraphs(doc):
    """Collect all paragraphs from body and tables."""
    items = []

    for i, para in enumerate(doc.paragraphs):
        items.append({"id": f"p{i}", "text": get_para_text(para), "para": para})

    # Deduplicate merged cells using the underlying XML element identity
    seen_tcs = set()
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                tc_id = id(cell._tc)
                if tc_id in seen_tcs:
                    continue
                seen_tcs.add(tc_id)
                for p_idx, para in enumerate(cell.paragraphs):
                    items.append(
                        {
                            "id": f"t{t_idx}r{r_idx}c{c_idx}p{p_idx}",
                            "text": get_para_text(para),
                            "para": para,
                        }
                    )

    return items


def call_claude(items_chunk):
    """Call Claude API to proofread a chunk of paragraphs."""
    non_empty = [(item["id"], item["text"]) for item in items_chunk if item["text"].strip()]
    if not non_empty:
        return {}

    text_block = "\n".join(f"[{pid}] {text}" for pid, text in non_empty)

    prompt = f"""あなたは日本語文書校正の専門家です。マンション管理組合の総会議案書を以下の方針で校正してください。

【校正方針】
1. 誤字・脱字の修正
2. 文法・語法の誤り修正
3. 敬語・文体の統一（丁寧語・尊敬語・謙譲語の適切な使用）
4. 管理組合・区分所有法の専門用語の正確な使用
5. 句読点の適切な配置
6. 数字・単位の統一（算用数字・漢数字の文脈に応じた使い分け）

【ルール】
- 意味・内容・文書の構成は一切変えずに、文章の誤りのみを修正すること
- 修正が必要な段落のみJSON形式で返すこと
- 修正不要な段落はJSONに含めないこと
- 修正後テキストは段落全文を返すこと（差分のみでなく）

【出力形式】以下のJSON形式のみで回答すること（前置き・説明文は不要）：
{{"corrections":[{{"id":"段落ID","corrected":"修正後テキスト全文","reason":"修正理由（30文字以内で簡潔に）"}}]}}

【校正対象】
{text_block}"""

    message = client.messages.create(
        model="claude-opus-4-7",
        max_tokens=8096,
        messages=[{"role": "user", "content": prompt}],
    )

    response_text = message.content[0].text

    json_match = re.search(r"\{.*\}", response_text, re.DOTALL)
    if not json_match:
        return {}
    try:
        result = json.loads(json_match.group())
        return {c["id"]: c for c in result.get("corrections", [])}
    except (json.JSONDecodeError, KeyError):
        return {}


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/correct", methods=["POST"])
def correct():
    if "file" not in request.files:
        return jsonify({"error": "ファイルが選択されていません"}), 400

    file = request.files["file"]
    if not file.filename.lower().endswith(".docx"):
        return jsonify({"error": ".docxファイルのみ対応しています"}), 400

    file_bytes = file.read()
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as e:
        return jsonify({"error": f"ファイル読み込みエラー: {str(e)}"}), 400

    items = collect_paragraphs(doc)
    id_to_item = {item["id"]: item for item in items}

    # Process paragraphs in chunks to stay within API limits
    all_corrections = {}
    for i in range(0, len(items), CHUNK_SIZE):
        chunk_corrections = call_claude(items[i : i + CHUNK_SIZE])
        all_corrections.update(chunk_corrections)

    changes = []
    for item_id, correction in all_corrections.items():
        item = id_to_item.get(item_id)
        if not item:
            continue
        original = item["text"]
        corrected = correction.get("corrected", "")
        if not corrected or original == corrected:
            continue
        # Sanity check: reject implausibly large rewrites
        if len(corrected) > 3 * max(len(original), 1) or len(corrected) < len(original) / 3:
            continue
        if apply_correction(item["para"], corrected):
            changes.append(
                {
                    "original": original,
                    "corrected": corrected,
                    "reason": correction.get("reason", ""),
                }
            )

    output = io.BytesIO()
    doc.save(output)
    doc_b64 = base64.b64encode(output.getvalue()).decode()

    base_name = re.sub(r"\.docx$", "", file.filename, flags=re.IGNORECASE)
    return jsonify(
        {
            "success": True,
            "correction_count": len(changes),
            "changes": changes,
            "document": doc_b64,
            "filename": f"{base_name}_校正済み.docx",
        }
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
